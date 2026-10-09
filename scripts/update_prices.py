import io
import json
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import requests
from openpyxl import load_workbook

URL = "https://www.meesman.nl/media/k5wd2q3j/handelskoersen-vanaf-2015-01-05.xlsx"
PRICE_FILE = Path("public/data/fund-prices.json")
DIVIDEND_FILE = Path("public/data/fund-dividends.json")

response = requests.get(URL, timeout=45)
response.raise_for_status()

workbook = load_workbook(
    io.BytesIO(response.content),
    read_only=True,
    data_only=True
)
sheet = workbook["Handelskoersen"]

if sheet.cell(4, 3).value != "Aandelen Wereldwijd Totaal A":
    raise ValueError("Kolom serie A niet gevonden; geen wijzigingen opgeslagen.")

values_by_date = defaultdict(list)

for row in sheet.iter_rows(min_row=5, values_only=True):
    date, value = row[0], row[2]

    if not isinstance(date, datetime):
        continue
    if not isinstance(value, (int, float)) or value <= 0:
        continue

    values_by_date[date.date().isoformat()].append(float(value))

prices = []
dividends = []

for date, values in sorted(values_by_date.items()):
    # Koersen liggen boven 10 euro; afzonderlijke
    # dividenduitkeringen liggen onder 10 euro.
    normal_prices = [v for v in values if v >= 10]
    payouts = [v for v in values if v < 10]

    if len(normal_prices) > 1 or len(payouts) > 1:
        raise ValueError(f"Onverwachte dubbele waarden op {date}")

    if payouts and not normal_prices:
        raise ValueError(f"Dividend zonder bijbehorende koers op {date}")

    if normal_prices:
        prices.append({
            "date": date,
            "a": round(normal_prices[0], 4)
        })

    if payouts:
        dividends.append({
            "date": date,
            "amount": round(payouts[0], 4)
        })

if len(prices) < 200:
    raise ValueError("Te weinig koersen; bestaande bestanden blijven behouden.")

if not dividends:
    raise ValueError("Geen dividendregels gevonden; update afgebroken.")

# Controleer of bestaande historische koersen behouden blijven.
if PRICE_FILE.exists():
    existing = json.loads(PRICE_FILE.read_text(encoding="utf-8"))
    old_prices = {
        item["date"]: float(item["a"])
        for item in existing.get("prices", [])
    }
    new_prices = {item["date"]: item["a"] for item in prices}

    for date, old_value in old_prices.items():
        if date not in new_prices:
            raise ValueError(f"Bestaande koersdatum ontbreekt: {date}")
        if abs(new_prices[date] - old_value) > 0.00011:
            raise ValueError(f"Bestaande koers gewijzigd: {date}")

price_data = {
    "fund": "Meesman Aandelen Wereldwijd Totaal",
    "source": URL,
    "note": "Handelskoersen serie A; dividenduitkeringen apart.",
    "prices": prices
}

dividend_data = {
    "fund": "Meesman Aandelen Wereldwijd Totaal",
    "series": "A",
    "source": URL,
    "dividends": dividends
}

PRICE_FILE.parent.mkdir(parents=True, exist_ok=True)

PRICE_FILE.write_text(
    json.dumps(price_data, ensure_ascii=False, indent=2) + "\n",
    encoding="utf-8"
)

DIVIDEND_FILE.write_text(
    json.dumps(dividend_data, ensure_ascii=False, indent=2) + "\n",
    encoding="utf-8"
)

print(f"{len(prices)} koersen verwerkt")
print(f"{len(dividends)} dividenduitkeringen verwerkt")
